"""Extract plain text from uploaded documents (PDF, DOCX, and text-like files)."""
import io
import os

# Cap how much extracted text we forward, to avoid blowing the context window.
MAX_CHARS = 60_000

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".yaml", ".yml",
    ".xml", ".html", ".htm", ".log", ".ini", ".cfg", ".toml", ".rtf",
    ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".c", ".h", ".cpp", ".hpp",
    ".cs", ".go", ".rs", ".rb", ".php", ".sh", ".bash", ".zsh", ".sql",
    ".css", ".scss", ".swift", ".kt", ".r", ".pl", ".lua", ".dart",
}


def _truncate(text: str) -> str:
    if len(text) > MAX_CHARS:
        return text[:MAX_CHARS] + "\n\n[... document truncated ...]"
    return text


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        if page_text.strip():
            parts.append(f"--- Page {i} ---\n{page_text.strip()}")
    return "\n\n".join(parts).strip()


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def extract_text(data: bytes, filename: str) -> str:
    """Return extracted text, or a note explaining why extraction failed."""
    ext = os.path.splitext(filename)[1].lower()
    try:
        if ext == ".pdf":
            text = _extract_pdf(data)
        elif ext in (".docx",):
            text = _extract_docx(data)
        elif ext in TEXT_EXTENSIONS or not ext:
            text = data.decode("utf-8", errors="replace")
        else:
            # Best-effort: try UTF-8 decode for unknown types.
            text = data.decode("utf-8", errors="replace")
    except Exception as exc:  # noqa: BLE001
        return f"[Could not extract text from this file: {exc}]"

    text = text.strip()
    if not text:
        return "[No extractable text found in this file.]"
    return _truncate(text)
